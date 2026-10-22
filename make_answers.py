import docx

doc = docx.Document()
doc.add_heading('MentorAI - Technical Methodology & Implementation Details', 0)
doc.add_paragraph('This document contains accurate technical details regarding the Machine Learning, Data Preprocessing, and Graph Network implementations used in the MentorAI project.')

qa = {
    '1. Which Models Were Used & Their Purpose?': '1. K-Means Clustering (Scikit-Learn): Used for Unsupervised Personality Archetype discovery. It segments the user base (students and mentors) into distinct behavioral clusters based purely on their 5-dimensional psychometric scores.\n2. Euclidean Distance Engine (K-Nearest Neighbors approach): Used as the core mathematical matching algorithm in the Allocation Engine to calculate compatibility scores between Mentors and Groups.\n3. Large Language Model (Llama 3.2 via Ollama): Used as the Generative AI component to analyze the structured trait data and output natural-language coaching briefings.',
    '2. Number of Clusters (k) and Selection Method': 'For the K-Means model, the optimal number of clusters (k) was initially evaluated using the Elbow Method (Within-Cluster Sum of Square - WCSS) and validated with Silhouette Scores on the baseline dataset, which indicated optimal groupings at k=3 and k=5. In the final live MentorAI application, the model is dynamic, allowing users to toggle k between 2 and 6 via the UI to observe macro and micro archetypes in real-time.',
    '3. Data Preprocessing Pipeline': 'For the original Big Data pipeline (the 400MB psychometric dataset):\n- Missing Data: Dropped rows with null values or incomplete test responses (zeroes).\n- Reverse-Scoring: Psychometric tests contain negatively keyed items (e.g., "I get stressed easily" vs "I am relaxed"). We mathematically reverse-scored these items before aggregation.\n- Normalization: Scaled raw Likert-scale responses (1-5) into standardized domain scores out of 100.\n- In the V2 live database, the preprocessing is handled on-the-fly when a user submits their 20-question assessment, directly parsing the JSON payload into aggregated Big Five dimensions.',
    '4. Network Analysis (Graph ML)': 'The Neural Web uses NetworkX and PyVis.\n- Edge Formation: Edges are formed using a Strict Similarity Threshold based on Euclidean distance in the 5D personality space. If the psychometric distance between two users is below the threshold (e.g., 1.5), a compatibility edge is drawn.\n- Topology: The screenshot showing ~64 nodes and 36-60 edges represents the seeded university population (Mentors + Students). \n- Physics & Communities: Instead of static centrality, we use a physics-based gravity simulation (forceAtlas2Based solver). Highly connected nodes (communities) are physically pulled together by gravity, naturally visualizing clusters without requiring explicit community-detection algorithms like Louvain.',
    '5. Tools & Big Data Stack': 'The project relies entirely on the optimized standard Python Scientific Stack to ensure low latency and easy deployment:\n- Pandas & NumPy for data manipulation and vectorized preprocessing.\n- Scikit-Learn for K-Means clustering and distance calculations.\n- NetworkX & PyVis for Graph ML topology and physics simulations.\n- Streamlit for the frontend UI.\n- PostgreSQL (psycopg2) for the relational database.\nWe did not use distributed computing frameworks like Spark or Dask, as Pandas chunking and optimized SQL queries were sufficient for the dataset size.',
}

for q, a in qa.items():
    doc.add_heading(q, level=2)
    doc.add_paragraph(a)

doc.save('MentorAI_Methodology_Answers.docx')
print('Word document saved successfully.')
